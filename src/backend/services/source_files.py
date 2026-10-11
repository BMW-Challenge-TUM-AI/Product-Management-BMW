"""Read source XLSX and PDF files while preserving their original locations."""
from __future__ import annotations

from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


SUPPORTED_SOURCE_FORMATS = {'.xlsx', '.pdf'}


def _json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def _columns(values: tuple[Any, ...] | list[Any]) -> list[str]:
    result: list[str] = []
    seen: dict[str, int] = {}
    for index, value in enumerate(values, start=1):
        base = str(value).strip() if value is not None else ''
        base = base or f'column_{index}'
        seen[base] = seen.get(base, 0) + 1
        result.append(base if seen[base] == 1 else f'{base}__{seen[base]}')
    return result


def read_source_file(file_path: str | Path) -> dict[str, Any]:
    """Read an XLSX workbook or text-based PDF, preserving sheet/row/page positions.

    XLSX results retain every row's values, including the header and blank rows,
    and also provide a convenient field mapping based on the first non-empty row.
    PDF results expose one text record per page. Scanned PDFs without an embedded
    text layer are reported with a warning; this reader does not run OCR.
    """
    path = Path(file_path)
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_SOURCE_FORMATS:
        raise ValueError('Supported source file formats are .xlsx and .pdf.')
    if not path.is_file():
        raise FileNotFoundError(f'Source file not found: {path.name}')

    if suffix == '.xlsx':
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheets = []
            for sheet in workbook.worksheets:
                rows = [
                    {'row': row_number, 'values': [_json_value(value) for value in values]}
                    for row_number, values in enumerate(sheet.iter_rows(values_only=True), start=1)
                ]
                header = next((row for row in rows if any(
                    value is not None and str(value).strip() for value in row['values'])), None)
                if header is None:
                    columns: list[str] = []
                    header_row = None
                else:
                    columns = _columns(header['values'])
                    header_row = header['row']
                records = []
                for row in rows:
                    if header_row is None or row['row'] <= header_row:
                        continue
                    if not any(value is not None and str(value).strip() for value in row['values']):
                        continue
                    fields = {
                        column: row['values'][index] if index < len(row['values']) else None
                        for index, column in enumerate(columns)
                    }
                    records.append({'row': row['row'], 'fields': fields, 'values': row['values']})
                sheets.append({
                    'name': sheet.title, 'header_row': header_row, 'columns': columns,
                    'rows': rows, 'records': records,
                })
            return {'file_name': path.name, 'format': 'xlsx', 'sheets': sheets}
        finally:
            workbook.close()

    import pymupdf

    document = pymupdf.open(path)
    try:
        pages = [{'page': number, 'text': page.get_text('text')}
                 for number, page in enumerate(document, start=1)]
    finally:
        document.close()
    return {
        'file_name': path.name,
        'format': 'pdf',
        'pages': pages,
        'warnings': ['No embedded text was found; scanned pages need a separate OCR step.']
        if not any(page['text'].strip() for page in pages) else [],
    }


def discover_source_files(folder: str | Path) -> list[dict[str, Any]]:
    """List supported files in a project folder without assuming BMW filenames."""
    root = Path(folder)
    if not root.is_dir():
        return []
    return [
        {'name': path.name, 'format': path.suffix.lower().lstrip('.'), 'available': True}
        for path in sorted(root.iterdir(), key=lambda item: item.name.casefold())
        if path.is_file() and path.suffix.lower() in SUPPORTED_SOURCE_FORMATS
    ]
