"""Synthetic adapter checks and optional local-file coverage checks."""
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook
from openpyxl import load_workbook

from src.backend.services import bmw_library


def workbook(path: Path, *, conflicting=False):
    book = Workbook()
    sheet = book.active
    sheet.title = 'Feedback_Explorer'
    sheet.append(['ID', 'Source', 'Brand', 'Sales Description', 'Country',
                  'Derivate (e-code)', 'Engine Type', 'Feedback Type',
                  'Customer Feedback', 'Vfc level2 3 Name', 'Vfc level2 Name', 'Vfc level3 Name'])
    sheet.append(['A1', 'survey', 'BMW', '1 Series', 'GB', 'F70', 'petrol',
                  'Likes', 'The seat is comfortable.', '', 'Comfort', 'Seat'])
    sheet.append(['A1', 'survey', 'BMW', '1 Series', 'GB', 'F70', 'petrol',
                  'Wants', 'The seat is uncomfortable.' if conflicting else 'The seat is comfortable.',
                  '', 'Comfort', 'Seat'])
    book.save(path)


class BmwLibraryTests(unittest.TestCase):
    def test_repeated_annotation_rows_become_one_traceable_comment(self):
        with tempfile.TemporaryDirectory() as temp, patch.object(bmw_library, 'RAW', Path(temp)):
            workbook(Path(temp) / bmw_library.FEEDBACK_FILES['F70'])
            records = bmw_library._feedback()[('F70', 'GB')]
            self.assertEqual(len(records), 1)
            self.assertEqual(records[0]['content'], 'The seat is comfortable.')
            self.assertEqual([a['row'] for a in records[0]['fields']['annotations']], [2, 3])
            self.assertEqual(records[0]['fields']['source_market_code'], 'GB')

    def test_conflicting_repeated_id_is_rejected(self):
        with tempfile.TemporaryDirectory() as temp, patch.object(bmw_library, 'RAW', Path(temp)):
            workbook(Path(temp) / bmw_library.FEEDBACK_FILES['F70'], conflicting=True)
            with self.assertRaisesRegex(ValueError, 'conflicting text'):
                bmw_library._feedback()

    def test_real_files_cover_distinct_scopes_without_relabeling(self):
        source_files = [*bmw_library.FEEDBACK_FILES.values(), *bmw_library.CATALOG_FILES.values(),
                        bmw_library.STUDY_FILE, bmw_library.SALES_FILE, bmw_library.EQUIPMENT_FILE]
        if not all((bmw_library.RAW / name).exists() for name in source_files):
            self.skipTest('Private BMW files are not present in ignored data/raw/.')
        source = bmw_library.manifest()
        by_scope = {(s['model_code'], s['market_code']): s for s in source['scopes']}
        self.assertGreater(len(source['scopes']), 9)
        self.assertEqual({s['series'] for s in source['scopes']}, {'1', '5', '7'})
        self.assertNotIn(('G60', 'CN'), by_scope)
        self.assertNotIn(('G68', 'EU'), by_scope)
        self.assertTrue({('F70', 'EU'), ('G60', 'EU'), ('G60', 'US'), ('G68', 'CN'),
                         ('G70', 'EU'), ('G70', 'US'), ('G70', 'CN')} <= set(by_scope))
        self.assertGreater(by_scope[('F70', 'EU')]['counts']['feedback'], 0)
        self.assertGreater(by_scope[('F70', 'EU')]['counts']['equipment'], 0)
        self.assertGreater(by_scope[('G68', 'CN')]['counts']['studies'], 0)
        self.assertGreater(by_scope[('G60', 'US')]['counts']['studies'], 0)
        eu = bmw_library.load_scope('F70', 'EU')
        original_markets = {r['fields']['source_market_code'] for d in eu['datasets'] for r in d['records']}
        self.assertIn('GB', original_markets)
        self.assertIn('DE', original_markets)
        self.assertIn('EU', original_markets)
        self.assertNotIn('JP', original_markets)
        self.assertTrue(all(r['fields']['analysis_market_code'] == 'EU'
                            for d in eu['datasets'] for r in d['records']))
        cn = bmw_library.load_scope('G68', 'CN')
        combined = [r for d in cn['datasets'] if d['kind'] == 'sales' for r in d['records']]
        self.assertTrue(combined and all(r['fields']['combined_model_scope'] for r in combined))
        self.assertTrue(any(r['fields']['future_value'] for r in combined))
        jp = bmw_library.load_scope('G60', 'JP')
        self.assertEqual(jp['market'], 'Japan')
        self.assertEqual(jp['counts']['feedback'], 53)
        self.assertEqual(jp['feedback_annotation_rows'], 76)
        with self.assertRaises(ValueError):
            bmw_library.load_scope('G60', 'CN')

    def test_feedback_counts_match_independent_workbook_id_and_row_counts(self):
        if not all((bmw_library.RAW / name).exists() for name in bmw_library.FEEDBACK_FILES.values()):
            self.skipTest('Private BMW feedback workbooks are not present in ignored data/raw/.')
        manifest = {(scope['model_code'], scope['market_code']): scope for scope in bmw_library.manifest()['scopes']}
        expected_totals = {'F70': (2128, 4076), 'G60': (4000, 5005), 'G70': (4000, 4690)}
        feedback = bmw_library._feedback()
        self.assertEqual(sum(len(records) for records in feedback.values()), 10128)
        self.assertEqual(sum(int(record['fields']['original_row_count'])
                             for records in feedback.values() for record in records), 13771)
        expected_analysis_scopes = {
            ('F70', 'EU'): (1433, 3010),
            ('G60', 'EU'): (111, 238),
            ('G60', 'US'): (3610, 4365),
            ('G70', 'EU'): (19, 44),
            ('G70', 'US'): (3792, 4418),
        }
        for key, expected in expected_analysis_scopes.items():
            self.assertEqual((manifest[key]['counts']['feedback'],
                              manifest[key]['feedback_annotation_rows']), expected)
        for model, filename in bmw_library.FEEDBACK_FILES.items():
            book = load_workbook(bmw_library.RAW / filename, read_only=True, data_only=True)
            sheet = book['Feedback_Explorer']
            rows = iter(sheet.values)
            headers = next(rows)
            unique_by_market = {}
            annotation_rows_by_market = {}
            for values in rows:
                row = dict(zip(headers, values))
                market = str(row['Country']).strip().upper()
                identity = str(row['ID']).strip()
                if not identity or not str(row['Customer Feedback'] or '').strip():
                    continue
                unique_by_market.setdefault(market, set()).add(identity)
                annotation_rows_by_market[market] = annotation_rows_by_market.get(market, 0) + 1
            book.close()
            unique_count = sum(map(len, unique_by_market.values()))
            annotation_count = sum(annotation_rows_by_market.values())
            self.assertEqual((unique_count, annotation_count), expected_totals[model])
            for analysis_scope, source_markets in (
                ('EU', bmw_library.EUROPE_SOURCE_CODES), ('US', {'US'}), ('CN', {'CN'}),
            ):
                expected_ids = sum(len(ids) for market, ids in unique_by_market.items()
                                   if market in source_markets)
                expected_rows = sum(rows for market, rows in annotation_rows_by_market.items()
                                    if market in source_markets)
                if expected_ids:
                    actual = manifest[(model, analysis_scope)]
                    self.assertEqual(actual['counts']['feedback'], expected_ids)
                    self.assertEqual(actual['feedback_annotation_rows'], expected_rows)
            for market, ids in unique_by_market.items():
                if market in bmw_library.ANALYSIS_MARKETS:
                    continue
                actual = manifest[(model, market)]
                self.assertEqual(actual['counts']['feedback'], len(ids))
                self.assertEqual(actual['feedback_annotation_rows'], annotation_rows_by_market[market])


if __name__ == '__main__':
    unittest.main()
